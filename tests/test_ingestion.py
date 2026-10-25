import io
import shutil

import docx
import pymupdf
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image, ImageDraw, ImageFont

from contractguard.config import get_settings
from contractguard.errors import IngestionError
from contractguard.ingestion.docx_loader import load_docx
from contractguard.ingestion.ocr import OcrResult, ocr_image, tesseract_available
from contractguard.ingestion.pdf_loader import load_pdf
from contractguard.ingestion.router import detect_file_type, load_document

HAS_TESSERACT = shutil.which("tesseract") is not None


# --------------------------------------------------------------------------- helpers


def make_pdf(path, pages):
    """Each page is a list of lines drawn as text."""
    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page()
        for i, line in enumerate(lines):
            page.insert_text((72, 72 + 16 * i), line, fontsize=11)
    doc.save(path)
    return path


def image_bytes(text="SCANNED CLAUSE TEXT", size=(1200, 300)):
    img = Image.new("L", size, 255)
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 48)
    except OSError:
        font = ImageFont.load_default()
    draw.text((40, 100), text, fill=0, font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue(), img


def make_mixed_pdf(path):
    """Page 1 has text; page 2 is an image only (a scanned page)."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "1. Definitions. This page has a real text layer.", fontsize=11)
    page = doc.new_page()
    png, _ = image_bytes()
    page.insert_image(pymupdf.Rect(36, 36, 560, 200), stream=png)
    doc.save(path)
    return path


def fake_ocr(confidence):
    def run(page, settings):
        return OcrResult(
            paragraphs=["2. Scanned clause text", "(a) a sub item"], confidence=confidence
        )

    return run


# --------------------------------------------------------------------------- router


def test_detect_file_type_accepts_real_files(tmp_path):
    pdf = make_pdf(tmp_path / "a.pdf", [["hello"]])
    d = docx.Document()
    d.add_paragraph("hello")
    d.save(tmp_path / "b.docx")
    assert detect_file_type(pdf) == "pdf"
    assert detect_file_type(tmp_path / "b.docx") == "docx"


def test_renamed_file_is_rejected(tmp_path):
    d = docx.Document()
    d.save(tmp_path / "real.docx")
    fake = tmp_path / "fake.pdf"
    fake.write_bytes((tmp_path / "real.docx").read_bytes())
    with pytest.raises(IngestionError) as err:
        detect_file_type(fake)
    assert err.value.code == IngestionError.UNSUPPORTED_FILE_TYPE


@pytest.mark.parametrize("name", ["contract.txt", "contract.doc", "contract"])
def test_unsupported_extension(tmp_path, name):
    f = tmp_path / name
    f.write_bytes(b"%PDF-1.7 pretending")
    with pytest.raises(IngestionError, match="Unsupported file type"):
        detect_file_type(f)


def test_file_size_limit(tmp_path, monkeypatch):
    pdf = make_pdf(tmp_path / "a.pdf", [["x" * 50]])
    settings = get_settings().model_copy(deep=True)
    settings.limits.max_file_mb = 0
    with pytest.raises(IngestionError) as err:
        load_document(pdf, settings)
    assert err.value.code == IngestionError.FILE_TOO_LARGE


def test_missing_file(tmp_path):
    with pytest.raises(IngestionError, match="not found"):
        load_document(tmp_path / "nope.pdf")


def test_empty_document_rejected(tmp_path):
    d = docx.Document()
    d.save(tmp_path / "empty.docx")
    with pytest.raises(IngestionError) as err:
        load_document(tmp_path / "empty.docx")
    assert err.value.code == IngestionError.EMPTY_DOCUMENT


# --------------------------------------------------------------------------- PDF


def test_text_pdf_blocks_keep_pages(tmp_path):
    pdf = make_pdf(tmp_path / "a.pdf", [["1. Term of the agreement"], ["2. Payment terms apply"]])
    doc = load_pdf(pdf)
    assert doc.num_pages == 2 and doc.ocr_pages == []
    assert [b.page for b in doc.blocks] == [1, 2]
    assert doc.blocks[0].text.startswith("1. Term")


def test_encrypted_pdf_rejected(tmp_path):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    path = tmp_path / "locked.pdf"
    doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o")
    with pytest.raises(IngestionError) as err:
        load_pdf(path)
    assert err.value.code == IngestionError.ENCRYPTED_PDF
    assert "Password-protected" in err.value.message


def test_corrupt_pdf_rejected(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.7\nthis is not a pdf")
    with pytest.raises(IngestionError) as err:
        load_pdf(path)
    assert err.value.code == IngestionError.CORRUPT_FILE


def test_page_limit(tmp_path):
    pdf = make_pdf(tmp_path / "a.pdf", [["page"]] * 3)
    settings = get_settings().model_copy(deep=True)
    settings.limits.max_pages = 2
    with pytest.raises(IngestionError) as err:
        load_pdf(pdf, settings)
    assert err.value.code == IngestionError.TOO_MANY_PAGES


def test_only_scanned_pages_go_to_ocr(tmp_path):
    doc = load_pdf(make_mixed_pdf(tmp_path / "mixed.pdf"), ocr_fn=fake_ocr(91.0))
    assert doc.ocr_pages == [2]
    assert [b.is_ocr for b in doc.blocks] == [False, True, True]
    assert doc.blocks[1].page == 2 and doc.blocks[1].ocr_confidence == 91.0
    assert doc.warnings == []


def test_low_ocr_confidence_warns(tmp_path):
    doc = load_pdf(make_mixed_pdf(tmp_path / "mixed.pdf"), ocr_fn=fake_ocr(52.0))
    assert doc.warnings == ["Page 2 OCR confidence 52% - verify manually"]


def test_text_page_with_image_is_not_scanned(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "A page with a logo and plenty of real text in its text layer.")
    png, _ = image_bytes("LOGO", (300, 100))
    page.insert_image(pymupdf.Rect(400, 700, 550, 750), stream=png)
    doc.save(tmp_path / "logo.pdf")

    def boom(page, settings):
        raise AssertionError("OCR must not run")

    assert load_pdf(tmp_path / "logo.pdf", ocr_fn=boom).ocr_pages == []


def test_missing_tesseract_gives_clear_error(tmp_path, monkeypatch):
    settings = get_settings().model_copy(deep=True)
    settings.ocr.tesseract_cmd = str(tmp_path / "no-such-tesseract")
    tesseract_available.cache_clear()
    with pytest.raises(IngestionError) as err:
        load_pdf(make_mixed_pdf(tmp_path / "mixed.pdf"), settings)
    tesseract_available.cache_clear()
    assert err.value.code == IngestionError.OCR_UNAVAILABLE
    assert "Tesseract" in err.value.message


@pytest.mark.skipif(not HAS_TESSERACT, reason="Tesseract not installed")
def test_real_tesseract_reads_image():
    _, img = image_bytes("TERMINATION NOTICE")
    result = ocr_image(img, get_settings().ocr)
    assert "TERMINATION" in " ".join(result.paragraphs)
    assert result.confidence is not None and result.confidence > 70


# --------------------------------------------------------------------------- DOCX


def _numbered_docx(path):
    d = docx.Document()
    numbering = d.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="77">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1.%2"/></w:lvl>'
        '<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>'
        '<w:lvlText w:val="(%3)"/></w:lvl>'
        "</w:abstractNum>"
    )
    first = numbering.find(qn("w:num"))
    if first is not None:
        first.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(
        parse_xml(f'<w:num {nsdecls("w")} w:numId="77"><w:abstractNumId w:val="77"/></w:num>')
    )

    def para(text, level, style=None):
        p = d.add_paragraph(text, style=style)
        ppr = p._p.get_or_add_pPr()
        num_pr = parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/><w:numId w:val="77"/></w:numPr>'
        )
        pstyle = ppr.find(qn("w:pStyle"))
        if pstyle is not None:
            pstyle.addnext(num_pr)
        else:
            ppr.insert(0, num_pr)

    para("Definitions", 0, "Heading 2")
    para("Term", 0, "Heading 2")
    para("The term is one year.", 1)
    para("Renewal requires agreement.", 1)
    para("first item", 2)
    para("Termination", 0, "Heading 2")
    para("Either party may terminate.", 1)
    d.add_paragraph("")  # empty paragraphs are skipped
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Service level", "Target"
    table.cell(1, 0).text, table.cell(1, 1).text = "Uptime", "99.5%"
    d.save(path)
    return path


def test_docx_numbering_is_rebuilt(tmp_path):
    doc = load_docx(_numbered_docx(tmp_path / "n.docx"))
    texts = [b.text for b in doc.blocks]
    assert texts[:7] == [
        "1. Definitions",
        "2. Term",
        "2.1 The term is one year.",
        "2.2 Renewal requires agreement.",
        "(a) first item",
        "3. Termination",
        "3.1 Either party may terminate.",  # level 1 restarts under a new level-0 item
    ]


def test_docx_tables_styles_and_no_pages(tmp_path):
    doc = load_docx(_numbered_docx(tmp_path / "n.docx"))
    assert doc.blocks[-2].text == "Service level | Target"
    assert doc.blocks[-1].text == "Uptime | 99.5%"
    assert doc.blocks[-1].style == "Table"
    assert doc.blocks[0].style == "Heading 2"
    assert all(b.page is None for b in doc.blocks)
    assert [b.para_index for b in doc.blocks] == list(range(len(doc.blocks)))


def test_docx_merged_cells_counted_once(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=1, cols=3)
    merged = t.cell(0, 0).merge(t.cell(0, 1))
    merged.text = "Wide"
    t.cell(0, 2).text = "Narrow"
    d.save(tmp_path / "m.docx")
    assert load_docx(tmp_path / "m.docx").blocks[0].text == "Wide | Narrow"


def test_docx_word_limit(tmp_path):
    d = docx.Document()
    d.add_paragraph("word " * 30)
    d.save(tmp_path / "w.docx")
    settings = get_settings().model_copy(deep=True)
    settings.limits.max_docx_words = 20
    with pytest.raises(IngestionError) as err:
        load_docx(tmp_path / "w.docx", settings)
    assert err.value.code == IngestionError.TOO_MANY_WORDS


def test_corrupt_docx(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"PK\x03\x04 not really a zip")
    with pytest.raises(IngestionError) as err:
        load_docx(bad)
    assert err.value.code == IngestionError.CORRUPT_FILE
