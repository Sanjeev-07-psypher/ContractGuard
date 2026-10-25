"""DOCX loading with Word auto-numbering rebuilt into the text (HLD 6.1).

Word keeps automatic numbers ("7.2") out of the paragraph text, in ``w:numPr``. We rebuild them
from numbering.xml with a counter per numId and level, so segmentation sees "7.2 ..." as a reader
does. Paragraphs and tables are read in document order; a table row becomes one block with cells
joined by " | ". DOCX has no pages: ``page`` stays None and locations use section numbers.
"""

from __future__ import annotations

import logging
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from contractguard.config import Settings, get_settings
from contractguard.errors import IngestionError
from contractguard.models import Block, Document

log = logging.getLogger(__name__)

_LEVEL_REF = re.compile(r"%(\d)")


@dataclass
class _Level:
    start: int = 1
    fmt: str = "decimal"
    text: str = "%1."


@dataclass
class Numbering:
    """numId -> level definitions, plus running counters (one list per numId)."""

    levels: dict[str, dict[int, _Level]] = field(default_factory=dict)
    counters: dict[str, list[int]] = field(default_factory=dict)

    @classmethod
    def from_document(cls, document: DocxDocument) -> Numbering:
        try:
            root = document.part.numbering_part.element
        except (KeyError, NotImplementedError):
            return cls()
        abstract: dict[str, dict[int, _Level]] = {}
        for an in root.findall(qn("w:abstractNum")):
            levels: dict[int, _Level] = {}
            for lvl in an.findall(qn("w:lvl")):
                levels[int(lvl.get(qn("w:ilvl")))] = _Level(
                    start=int(_val(lvl, "w:start") or 1),
                    fmt=_val(lvl, "w:numFmt") or "decimal",
                    text=_val(lvl, "w:lvlText") or "",
                )
            abstract[an.get(qn("w:abstractNumId"))] = levels
        num = cls()
        for n in root.findall(qn("w:num")):
            ref = n.find(qn("w:abstractNumId"))
            if ref is not None:
                num.levels[n.get(qn("w:numId"))] = abstract.get(ref.get(qn("w:val")), {})
        return num

    def label(self, num_id: str, ilvl: int) -> str | None:
        """Advance the counter for (num_id, ilvl) and return the rendered label, e.g. "7.2"."""
        levels = self.levels.get(num_id)
        if not levels or ilvl not in levels:
            return None
        level = levels[ilvl]
        if level.fmt in ("bullet", "none"):
            return None
        counters = self.counters.setdefault(num_id, [])
        while len(counters) <= ilvl:
            lvl = levels.get(len(counters), _Level())
            counters.append(lvl.start - 1)
        counters[ilvl] += 1
        for deeper in range(ilvl + 1, len(counters)):  # restart lower levels
            counters[deeper] = levels.get(deeper, _Level()).start - 1

        def render(match: re.Match[str]) -> str:
            k = int(match.group(1)) - 1
            value = counters[k] if k < len(counters) else 0
            return _format(value, levels.get(k, _Level()).fmt)

        return _LEVEL_REF.sub(render, level.text).strip() or None


def _val(el, tag: str) -> str | None:
    child = el.find(qn(tag))
    return None if child is None else child.get(qn("w:val"))


def _roman(n: int) -> str:
    pairs = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"),
             (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]  # fmt: skip
    out = ""
    for value, sym in pairs:
        while n >= value:
            out += sym
            n -= value
    return out


def _format(n: int, fmt: str) -> str:
    if fmt == "lowerLetter":
        return chr(ord("a") + (n - 1) % 26)
    if fmt == "upperLetter":
        return chr(ord("A") + (n - 1) % 26)
    if fmt == "lowerRoman":
        return _roman(n)
    if fmt == "upperRoman":
        return _roman(n).upper()
    return str(n)


def _num_pr(paragraph: Paragraph) -> tuple[str, int] | None:
    """(numId, ilvl) from the paragraph, else from its style; None if not numbered."""
    candidates = [paragraph._p.pPr]
    style = paragraph.style
    while style is not None:
        candidates.append(style.element.pPr)
        style = style.base_style
    for ppr in candidates:
        num_pr = None if ppr is None else ppr.find(qn("w:numPr"))
        if num_pr is None:
            continue
        num_id = _val(num_pr, "w:numId")
        if num_id is None or num_id == "0":  # numId 0 switches numbering off
            return None
        return num_id, int(_val(num_pr, "w:ilvl") or 0)
    return None


def _row_text(row) -> str:
    cells, seen = [], set()
    for cell in row.cells:  # merged cells repeat the same element; keep one
        if id(cell._tc) in seen:
            continue
        seen.add(id(cell._tc))
        text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
        cells.append(text)
    return " | ".join(cells).strip(" |")


def load_docx(path: Path, settings: Settings | None = None) -> Document:
    settings = settings or get_settings()
    try:
        document = docx.Document(str(path))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError, ValueError) as exc:
        raise IngestionError(
            "The DOCX file could not be opened; it may be corrupt.", IngestionError.CORRUPT_FILE
        ) from exc

    numbering = Numbering.from_document(document)
    blocks: list[Block] = []
    words = 0

    def add(text: str, style: str | None) -> None:
        nonlocal words
        blocks.append(Block(text=text, para_index=len(blocks), style=style))
        words += len(text.split())
        if words > settings.limits.max_docx_words:
            raise IngestionError(
                f"The document has more than {settings.limits.max_docx_words} words.",
                IngestionError.TOO_MANY_WORDS,
            )

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            num = _num_pr(paragraph)
            label = numbering.label(*num) if num else None
            style = paragraph.style.name if paragraph.style is not None else None
            add(f"{label} {text}" if label else text, style)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, document).rows:
                text = _row_text(row)
                if text:
                    add(text, "Table")
    log.info("docx %s: %d blocks, %d words", path.name, len(blocks), words)
    return Document(source_path=str(path), file_type="docx", num_pages=None, blocks=blocks)
