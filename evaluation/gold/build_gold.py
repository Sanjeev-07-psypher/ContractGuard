"""Build the synthetic gold set from its YAML sources (HLD 12.2).

For every ``evaluation/gold/<id>/contract.yaml`` this writes, next to it:

- ``contract.docx``         the original (python-docx)
- ``contract_text.pdf``     a text PDF exported from the DOCX by LibreOffice (headless)
- ``contract_scanned.pdf``  an image-only PDF: the text PDF rendered, skewed, blurred and noised
- ``gold.json``             clause boundaries, categories, facts and expected rule statuses

Usage:
    python evaluation/gold/build_gold.py            # all contracts
    python evaluation/gold/build_gold.py gc01 gc03  # selected contracts
    python evaluation/gold/build_gold.py --no-pdf   # DOCX + gold.json only (no LibreOffice)

Content is reproducible for a given source and tool versions (seeded scan noise); PDF bytes differ
between runs because LibreOffice embeds timestamps.
"""

from __future__ import annotations

import argparse
import io
import json
import random
import shutil
import subprocess
import sys
import tempfile
import zlib
from pathlib import Path
from typing import Any

import docx
import numpy as np
import pymupdf
import yaml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor
from PIL import Image, ImageFilter

GOLD_DIR = Path(__file__).resolve().parent
REPO_ROOT = GOLD_DIR.parents[1]
CATEGORIES = {
    "CONFIDENTIALITY", "TERMINATION", "PAYMENT", "LIABILITY", "INDEMNIFICATION", "IP",
    "GOVERNING_LAW", "RENEWAL", "SLA", "OTHER",
}  # fmt: skip
STATUSES = {"COMPLIANT", "NON_COMPLIANT", "MISSING", "NEEDS_REVIEW", "NOT_APPLICABLE"}
FACT_KEYS = {
    "termination.notice_days", "termination.counterparty_unilateral", "liability.has_cap",
    "liability.is_uncapped", "renewal.is_automatic", "payment.due_days",
    "governing_law.jurisdiction",
}  # fmt: skip
NUM_ID = 90  # numbering id used for Word auto-numbered contracts
SCAN_DPI = 200
FIRST_WORDS = 8

_NORMALIZE = str.maketrans({"–": "-", "—": "-", "“": '"', "”": '"',
                            "‘": "'", "’": "'", " ": " "})  # fmt: skip


# --------------------------------------------------------------------------- source model


def heading_text(spec: dict, clause: dict) -> str:
    """The heading line as a reader sees it (number included)."""
    n, h = clause["number"], clause["heading"]
    style = spec["heading_style"]
    if style == "num_title":
        return f"{n}. {h}"
    if style == "num_caps":
        return f"{n}. {h.upper()}"
    if style == "article":
        return f"ARTICLE {n} – {h.upper()}"
    raise ValueError(f"unknown heading_style {style!r}")


def units(spec: dict) -> list[dict]:
    """Flatten the source into gold clause units in reading order.

    A numbered heading with sub-clauses and no body of its own is not a unit: its heading line
    belongs to the first sub-clause (HLD 6.3 length guard).
    """
    out: list[dict] = [
        {"number": None, "heading": None, "kind": "PREAMBLE", "categories": [],
         "section_path": "PREAMBLE", "lines": [spec["title"], *spec["preamble"]]}
    ]  # fmt: skip
    for c in spec["clauses"]:
        head = heading_text(spec, c)
        kind = c.get("kind", "BODY")
        if c.get("sub"):
            if c.get("body"):
                raise ValueError(f"{spec['id']} clause {c['number']}: body and sub both set")
            for i, s in enumerate(c["sub"]):
                lines = [f"{s['number']} {s['body'][0]}", *s["body"][1:]]
                out.append({
                    "number": s["number"], "heading": c["heading"] if i == 0 else None,
                    "kind": s.get("kind", kind), "categories": s["categories"],
                    "section_path": f"{c['number']} > {s['number']}",
                    "lines": [head, *lines] if i == 0 else lines,
                })  # fmt: skip
        else:
            out.append({
                "number": c["number"], "heading": c["heading"], "kind": kind,
                "categories": c["categories"], "section_path": c["number"],
                "lines": [head, *c["body"]],
            })  # fmt: skip
    out.append({"number": None, "heading": None, "kind": "SIGNATURE", "categories": [],
                "section_path": "SIGNATURE", "lines": spec["signature"]})  # fmt: skip
    for e in spec.get("exhibits", []):
        rows = [" | ".join(r) for r in e.get("table", [])]
        out.append({"number": None, "heading": e["heading"], "kind": "EXHIBIT",
                    "categories": e["categories"], "section_path": e["heading"],
                    "lines": [e["heading"], *e.get("body", []), *rows]})  # fmt: skip
    return out


def validate(spec: dict, policy: dict) -> None:
    sid = spec["id"]
    us = units(spec)
    numbers = {u["number"] for u in us if u["number"]}
    numbers |= {c["number"] for c in spec["clauses"]}
    for u in us:
        bad = set(u["categories"]) - CATEGORIES
        if bad:
            raise ValueError(f"{sid} {u['section_path']}: unknown categories {bad}")
        if u["kind"] in ("BODY", "EXHIBIT") and not u["categories"]:
            raise ValueError(f"{sid} {u['section_path']}: BODY/EXHIBIT clause needs categories")
    gold = spec["gold"]
    if set(gold["facts"]) != FACT_KEYS:
        raise ValueError(f"{sid}: facts must have exactly {sorted(FACT_KEYS)}")
    rules = {r["id"]: r for r in policy["rules"]}
    if set(gold["rules"]) != set(rules):
        raise ValueError(f"{sid}: gold rules must cover {sorted(rules)}")
    for rid, g in gold["rules"].items():
        if g["status"] not in STATUSES:
            raise ValueError(f"{sid} {rid}: bad status {g['status']}")
        applies = spec["contract_type"] in rules[rid]["applies_to"]
        if not applies and g["status"] != "NOT_APPLICABLE":
            raise ValueError(f"{sid} {rid}: rule does not apply to {spec['contract_type']}")
        if g["status"] in ("MISSING", "NOT_APPLICABLE") and g.get("clauses"):
            raise ValueError(f"{sid} {rid}: {g['status']} must not cite clauses")
        for n in g.get("clauses", []):
            if n not in numbers:
                raise ValueError(f"{sid} {rid}: cites unknown clause {n}")


# --------------------------------------------------------------------------- DOCX


def _add_numbering(document: docx.document.Document) -> None:
    numbering = document.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{NUM_ID}">'
        '<w:multiLevelType w:val="multilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:ind w:left="431" w:hanging="431"/></w:pPr></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1.%2"/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:ind w:left="862" w:hanging="578"/></w:pPr></w:lvl>'
        "</w:abstractNum>"
    )
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        first_num.addprevious(abstract)
    numbering.append(
        parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{NUM_ID}"><w:abstractNumId w:val="{NUM_ID}"/></w:num>'
        )
    )


def _set_num(paragraph, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/><w:numId w:val="{NUM_ID}"/></w:numPr>'
    )
    style = ppr.find(qn("w:pStyle"))
    if style is not None:
        style.addnext(num_pr)
    else:
        ppr.insert(0, num_pr)


def _field(paragraph, instr: str) -> None:
    paragraph._p.append(
        parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="{instr}"><w:r><w:t>1</w:t></w:r></w:fldSimple>'
        )  # fmt: skip
    )


def _heading(document, text: str):
    p = document.add_paragraph(style="Heading 2")
    run = p.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.size = Pt(11)
    return p


def build_docx(spec: dict, path: Path) -> None:
    auto = spec["numbering"] == "auto"
    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = "Liberation Serif"
    normal.font.size = Pt(11)
    if auto:
        _add_numbering(d)

    section = d.sections[0]
    header = section.header.paragraphs[0]
    header.text = f"{spec['short_title']} | {spec['our_org']} and {spec['counterparty']}"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    _field(footer, "PAGE")
    footer.add_run(" of ")
    _field(footer, "NUMPAGES")

    title = d.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(spec["title"]).bold = True
    for line in spec["preamble"]:
        d.add_paragraph(line)

    for c in spec["clauses"]:
        if auto:
            p = _heading(d, c["heading"])
            _set_num(p, 0)
        else:
            _heading(d, heading_text(spec, c))
        for line in c.get("body", []):
            d.add_paragraph(line)
        for s in c.get("sub", []):
            if auto:
                _set_num(d.add_paragraph(s["body"][0]), 1)
            else:
                d.add_paragraph(f"{s['number']} {s['body'][0]}")
            for line in s["body"][1:]:
                d.add_paragraph(line)

    for line in spec["signature"]:
        d.add_paragraph(line)
    for e in spec.get("exhibits", []):
        _heading(d, e["heading"])
        for line in e.get("body", []):
            d.add_paragraph(line)
        if e.get("table"):
            rows = e["table"]
            table = d.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for col, val in enumerate(row):
                    table.cell(r, col).text = val
    d.save(path)


# --------------------------------------------------------------------------- PDFs


def docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) not found; install it or use --no-pdf")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless",
             "--convert-to", "pdf", "--outdir", tmp, str(docx_path)],
            check=True, capture_output=True, timeout=180,
        )  # fmt: skip
        shutil.move(Path(tmp) / f"{docx_path.stem}.pdf", pdf_path)


def scan_pdf(text_pdf: Path, out_pdf: Path, seed_key: str) -> None:
    """Simulate a print-and-scan: grayscale raster, slight skew, blur, noise, JPEG."""
    src = pymupdf.open(text_pdf)
    dst = pymupdf.open()
    for i, page in enumerate(src):
        seed = zlib.crc32(f"{seed_key}-{i}".encode())
        rnd = random.Random(seed)
        rng = np.random.default_rng(seed)
        pix = page.get_pixmap(dpi=SCAN_DPI, colorspace=pymupdf.csGRAY)
        img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
        img = img.rotate(rnd.uniform(-0.8, 0.8), resample=Image.BICUBIC, fillcolor=255)
        arr = np.asarray(img, dtype=np.float32)
        arr = arr * 0.9 + 12 + rng.normal(0, 6, arr.shape)  # duller paper, sensor noise
        specks = rng.random(arr.shape) < 0.0003
        arr[specks] = rng.uniform(0, 120, specks.sum())
        img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
        img = img.filter(ImageFilter.GaussianBlur(radius=0.7))  # also smooths noise for JPEG
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=60)
        out_page = dst.new_page(width=page.rect.width, height=page.rect.height)
        out_page.insert_image(out_page.rect, stream=buf.getvalue())
    dst.save(out_pdf, deflate=True, garbage=4)


# --------------------------------------------------------------------------- gold.json


def gold_json(spec: dict) -> dict[str, Any]:
    clauses = []
    for i, u in enumerate(units(spec), start=1):
        text = " ".join(u["lines"]).translate(_NORMALIZE)
        clauses.append({
            "index": i, "number": u["number"], "heading": u["heading"], "kind": u["kind"],
            "section_path": u["section_path"],
            "first_words": " ".join(text.split()[:FIRST_WORDS]),
            "categories": u["categories"],
        })  # fmt: skip
    g = spec["gold"]
    return {
        "id": spec["id"],
        "title": spec["title"],
        "contract_type": spec["contract_type"],
        "our_org": spec["our_org"],
        "files": {
            "docx": "contract.docx",
            "text_pdf": "contract_text.pdf",
            "scanned_pdf": "contract_scanned.pdf",
        },  # fmt: skip
        "annotation": {
            "status": "DRAFT",
            "drafted_by": "Claude (generated from contract.yaml)",
            "labeller_1": None,
            "labeller_2": None,
            "adjudicated": False,
            "notes": "Needs two independent human labellers per GUIDELINES.md before use.",
        },
        "meta": g["meta"],
        "clauses": clauses,
        "facts": g["facts"],
        "rules": {
            rid: {"status": r["status"], "clauses": r.get("clauses", [])}
            for rid, r in sorted(g["rules"].items())
        },  # fmt: skip
    }


# --------------------------------------------------------------------------- main


def build(contract_dir: Path, policy: dict, make_pdf: bool) -> None:
    spec = yaml.safe_load((contract_dir / "contract.yaml").read_text(encoding="utf-8"))
    validate(spec, policy)
    build_docx(spec, contract_dir / "contract.docx")
    if make_pdf:
        docx_to_pdf(contract_dir / "contract.docx", contract_dir / "contract_text.pdf")
        scan_pdf(contract_dir / "contract_text.pdf", contract_dir / "contract_scanned.pdf",
                 spec["id"])  # fmt: skip
    out = json.dumps(gold_json(spec), indent=2, ensure_ascii=False)
    (contract_dir / "gold.json").write_text(out + "\n", encoding="utf-8")
    print(f"built {spec['id']}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build gold contracts from YAML sources.")
    parser.add_argument("ids", nargs="*", help="contract ids (default: all)")
    parser.add_argument("--no-pdf", action="store_true", help="skip LibreOffice and scanning")
    args = parser.parse_args(argv)
    policy = json.loads((REPO_ROOT / "policies" / "default_vendor_v1.json").read_text())
    dirs = sorted(p.parent for p in GOLD_DIR.glob("*/contract.yaml"))
    if args.ids:
        dirs = [d for d in dirs if d.name in args.ids]
    for d in dirs:
        build(d, policy, make_pdf=not args.no_pdf)
    return 0


if __name__ == "__main__":
    sys.exit(main())
